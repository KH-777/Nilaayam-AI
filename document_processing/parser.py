"""Document parser for Nilaayam AI.

Extracts text, metadata, and structural layout from PDF, DOCX, and image documents.
Automatically detects scanned PDFs and routes pages through the OCR engine.
"""

import logging
import os
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

# Handle import errors gracefully with try/except
try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import docx
except ImportError:
    docx = None

try:
    from PIL import Image
except ImportError:
    Image = None

from document_processing.ocr_engine import (
    extract_text_from_image,
    extract_text_from_pdf_page,
    is_tesseract_available,
)

logger = logging.getLogger(__name__)


@dataclass
class PageResult:
    """Extraction result for an individual document page."""

    page_num: int
    text: str
    confidence: float  # 1.0 for native text, 0.0-1.0 for OCR
    is_ocr: bool
    width: float
    height: float


@dataclass
class ParseResult:
    """Aggregated extraction result for the complete document."""

    full_text: str
    pages: List[PageResult]
    doc_type_hint: str  # 'rental_agreement'|'patta'|'ec'|'sale_deed'|'unknown'
    page_count: int
    ocr_applied: bool
    overall_confidence: float
    warnings: List[str] = field(default_factory=list)
    file_format: str = ""  # 'pdf'|'docx'|'image'


def _guess_doc_type(text: str) -> str:
    """Guess the document type from text content using keyword indicators.

    Matches English and Tamil keywords specific to Tamil Nadu real estate documents.

    Args:
        text: The extracted text to analyze.

    Returns:
        One of 'rental_agreement', 'patta', 'ec', 'sale_deed', or 'unknown'.
    """
    if not text or not text.strip():
        return "unknown"

    lower_text = text.lower()

    # Document type keyword patterns
    rental_keywords = [
        "lease agreement",
        "rental agreement",
        "tenancy agreement",
        "tenant",
        "landlord",
        "monthly rent",
        "security deposit",
        "lessor",
        "lessee",
        "குத்தகை",
        "வாடகை",
        "வாடகை ஒப்பந்தம்",
    ]

    patta_keywords = [
        "patta",
        "பட்டா",
        "chitta",
        "சிட்டா",
        "survey number",
        "புல எண்",
        "taluk",
        "வட்டம்",
        "village",
        "கிராமம்",
        "revenue department",
        "வருவாய்த்துறை",
    ]

    ec_keywords = [
        "encumbrance certificate",
        "ec",
        "வில்லங்க சான்று",
        "வில்லங்கம்",
        "transactions",
        "mortgage",
        "sub-registrar",
        "nil encumbrance",
        "சுமை விவரம்",
    ]

    sale_deed_keywords = [
        "sale deed",
        "விற்பனை பத்திரம்",
        "கிரய பத்திரம்",
        "vendor",
        "vendee",
        "sale consideration",
        "purchaser",
        "absolute sale",
        "registered",
    ]

    scores: Dict[str, int] = {
        "rental_agreement": 0,
        "patta": 0,
        "ec": 0,
        "sale_deed": 0,
    }

    # Score each type based on keyword occurrences
    for kw in rental_keywords:
        if kw in lower_text:
            scores["rental_agreement"] += 2 if len(kw.split()) > 1 else 1

    for kw in patta_keywords:
        if kw in lower_text:
            scores["patta"] += 2 if len(kw.split()) > 1 else 1

    for kw in ec_keywords:
        if kw in lower_text:
            scores["ec"] += 2 if len(kw.split()) > 1 else 1

    for kw in sale_deed_keywords:
        if kw in lower_text:
            scores["sale_deed"] += 2 if len(kw.split()) > 1 else 1

    best_type, max_score = max(scores.items(), key=lambda item: item[1])
    if max_score >= 2:
        return best_type

    return "unknown"


def _is_pdf_scanned(file_path: str) -> bool:
    """Check if a PDF file is primarily scanned/image-based rather than selectable text.

    Uses PyMuPDF (fitz) or pdfplumber to check average characters per page.
    A PDF is considered scanned if the average character count per page is under 20.

    Args:
        file_path: Path to the PDF file.

    Returns:
        bool: True if scanned/image-based, False if digital text is selectable.
    """
    # 1. Try PyMuPDF (fastest)
    if fitz is not None:
        try:
            doc = fitz.open(file_path)
            total_pages = len(doc)
            if total_pages == 0:
                doc.close()
                return True

            total_chars = 0
            for page in doc:
                text = page.get_text().strip()
                total_chars += len(text)
            doc.close()

            avg_chars = total_chars / total_pages
            return avg_chars < 20
        except Exception as exc:
            logger.warning("fitz failed to determine if PDF is scanned: %s", exc)

    # 2. Try pdfplumber fallback
    if pdfplumber is not None:
        try:
            with pdfplumber.open(file_path) as pdf:
                total_pages = len(pdf.pages)
                if total_pages == 0:
                    return True
                total_chars = sum(len((p.extract_text() or "").strip()) for p in pdf.pages)
                avg_chars = total_chars / total_pages
                return avg_chars < 20
        except Exception as exc:
            logger.warning("pdfplumber failed to determine if PDF is scanned: %s", exc)

    # Default assumption if libraries not available or errors occur
    return False


def _extract_pdf_text(file_path: str) -> List[PageResult]:
    """Extract selectable text and dimensions from a native PDF using pdfplumber or fitz.

    Args:
        file_path: Path to the PDF file.

    Returns:
        List[PageResult]: Extracted pages with confidence 1.0 for valid text.
    """
    results: List[PageResult] = []

    # 1. Primary extractor: pdfplumber (best for layout and table text)
    if pdfplumber is not None:
        try:
            with pdfplumber.open(file_path) as pdf:
                for idx, page in enumerate(pdf.pages, start=1):
                    raw_text = page.extract_text() or ""
                    clean_text = raw_text.strip()
                    width = float(page.width or 595.0)
                    height = float(page.height or 842.0)

                    # Mark empty pages with 0.0 confidence
                    confidence = 1.0 if clean_text else 0.0
                    results.append(
                        PageResult(
                            page_num=idx,
                            text=clean_text,
                            confidence=confidence,
                            is_ocr=False,
                            width=width,
                            height=height,
                        )
                    )
            return results
        except Exception as exc:
            logger.warning("pdfplumber extraction failed, falling back to fitz: %s", exc)

    # 2. Fallback extractor: PyMuPDF (fitz)
    if fitz is not None:
        try:
            doc = fitz.open(file_path)
            for idx, page in enumerate(doc, start=1):
                clean_text = page.get_text().strip()
                rect = page.rect
                confidence = 1.0 if clean_text else 0.0
                results.append(
                    PageResult(
                        page_num=idx,
                        text=clean_text,
                        confidence=confidence,
                        is_ocr=False,
                        width=float(rect.width),
                        height=float(rect.height),
                    )
                )
            doc.close()
            return results
        except Exception as exc:
            logger.error("fitz extraction failed: %s", exc)

    logger.error("Neither pdfplumber nor fitz could extract text from '%s'", file_path)
    return [
        PageResult(
            page_num=1,
            text="",
            confidence=0.0,
            is_ocr=False,
            width=0.0,
            height=0.0,
        )
    ]


def _extract_docx_text(file_path: str) -> List[PageResult]:
    """Extract text from a Word DOCX file using python-docx.

    Estimates pages by grouping every 40 non-empty paragraphs into 1 virtual page.

    Args:
        file_path: Path to the .docx file.

    Returns:
        List[PageResult]: Virtual pages extracted from DOCX.
    """
    if docx is None:
        logger.error("python-docx is not installed; cannot extract DOCX text.")
        return [
            PageResult(
                page_num=1,
                text="",
                confidence=0.0,
                is_ocr=False,
                width=612.0,
                height=792.0,
            )
        ]

    try:
        doc = docx.Document(file_path)
        paragraphs: List[str] = []

        # Extract regular paragraphs
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                paragraphs.append(t)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))

        if not paragraphs:
            return [
                PageResult(
                    page_num=1,
                    text="",
                    confidence=1.0,
                    is_ocr=False,
                    width=612.0,
                    height=792.0,
                )
            ]

        # Group every 40 paragraphs into 1 page
        chunk_size = 40
        pages: List[PageResult] = []
        for i in range(0, len(paragraphs), chunk_size):
            chunk = paragraphs[i : i + chunk_size]
            page_num = (i // chunk_size) + 1
            page_text = "\n\n".join(chunk)
            pages.append(
                PageResult(
                    page_num=page_num,
                    text=page_text,
                    confidence=1.0,
                    is_ocr=False,
                    width=612.0,
                    height=792.0,
                )
            )

        return pages

    except Exception as exc:
        logger.error("Failed to extract text from DOCX '%s': %s", file_path, exc)
        return [
            PageResult(
                page_num=1,
                text="",
                confidence=0.0,
                is_ocr=False,
                width=612.0,
                height=792.0,
            )
        ]


def _extract_image_text(file_path: str) -> List[PageResult]:
    """Extract text from an image file using the OCR engine.

    Args:
        file_path: Path to the image file.

    Returns:
        List[PageResult]: Single PageResult containing OCR output.
    """
    width, height = 800.0, 1100.0
    if Image is not None:
        try:
            with Image.open(file_path) as img:
                width = float(img.width)
                height = float(img.height)
        except Exception as exc:
            logger.debug("Could not read image dimensions: %s", exc)

    ocr_result = extract_text_from_image(file_path)
    text = ocr_result.get("text", "")
    confidence = float(ocr_result.get("confidence", 0.0))

    return [
        PageResult(
            page_num=1,
            text=text,
            confidence=confidence,
            is_ocr=True,
            width=width,
            height=height,
        )
    ]


def parse_document(
    file_path: str,
    file_content: Optional[bytes] = None,
) -> ParseResult:
    """Parse a document from a file path or in-memory raw bytes.

    Automatically detects file type, checks whether PDFs are scanned, routes to OCR
    when appropriate, and constructs a complete ParseResult.

    Args:
        file_path: Original file path or filename.
        file_content: Optional raw bytes of the file.

    Returns:
        ParseResult containing all extracted pages, confidence metrics, and doc type hint.
    """
    temp_file_created = False
    actual_path = file_path
    warnings: List[str] = []

    # If raw bytes provided, write to a temporary file
    if file_content is not None:
        suffix = Path(file_path).suffix if file_path else ".pdf"
        if not suffix:
            suffix = ".pdf"
        try:
            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                tmp.write(file_content)
                actual_path = tmp.name
                temp_file_created = True
        except Exception as exc:
            logger.error("Failed to write temporary file for parsing: %s", exc)
            return ParseResult(
                full_text="",
                pages=[],
                doc_type_hint="unknown",
                page_count=0,
                ocr_applied=False,
                overall_confidence=0.0,
                warnings=[f"Failed to write temporary document bytes: {exc}"],
                file_format="",
            )

    try:
        ext = Path(actual_path).suffix.lower()
        file_format = ""
        pages: List[PageResult] = []
        ocr_applied = False

        # --- 1. PDF Documents ---
        if ext == ".pdf":
            file_format = "pdf"
            is_scanned = _is_pdf_scanned(actual_path)

            if is_scanned:
                logger.info("PDF '%s' is scanned/image-based; running OCR.", file_path)
                ocr_applied = True

                # Determine total pages
                total_pages = 1
                doc_fitz = None
                if fitz is not None:
                    try:
                        doc_fitz = fitz.open(actual_path)
                        total_pages = len(doc_fitz)
                    except Exception as exc:
                        logger.warning("Could not open PDF with fitz to count pages: %s", exc)

                for p_no in range(1, total_pages + 1):
                    ocr_res = extract_text_from_pdf_page(actual_path, p_no)
                    p_width, p_height = 595.0, 842.0

                    if doc_fitz is not None and p_no <= len(doc_fitz):
                        try:
                            rect = doc_fitz[p_no - 1].rect
                            p_width = float(rect.width)
                            p_height = float(rect.height)
                        except Exception:
                            pass

                    p_text = ocr_res.get("text", "")
                    p_conf = float(ocr_res.get("confidence", 0.0))
                    pages.append(
                        PageResult(
                            page_num=p_no,
                            text=p_text,
                            confidence=p_conf,
                            is_ocr=True,
                            width=p_width,
                            height=p_height,
                        )
                    )

                if doc_fitz is not None:
                    try:
                        doc_fitz.close()
                    except Exception:
                        pass
            else:
                logger.info("PDF '%s' contains selectable digital text.", file_path)
                pages = _extract_pdf_text(actual_path)
                ocr_applied = False

        # --- 2. Word Documents ---
        elif ext in (".docx", ".doc"):
            file_format = "docx"
            if ext == ".doc":
                warnings.append("Legacy .doc format detected; conversion to .docx is recommended.")
            pages = _extract_docx_text(actual_path)
            ocr_applied = False

        # --- 3. Image Files ---
        elif ext in (".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".webp"):
            file_format = "image"
            ocr_applied = True
            pages = _extract_image_text(actual_path)

        # --- 4. Unsupported Format ---
        else:
            file_format = ext.lstrip(".") or "unknown"
            warnings.append(f"Unsupported file format: '{ext}'. Extraction may be incomplete.")
            pages = [
                PageResult(
                    page_num=1,
                    text="",
                    confidence=0.0,
                    is_ocr=False,
                    width=0.0,
                    height=0.0,
                )
            ]
            ocr_applied = False

        # Ensure we have at least one PageResult
        if not pages:
            pages = [
                PageResult(
                    page_num=1,
                    text="",
                    confidence=0.0,
                    is_ocr=ocr_applied,
                    width=0.0,
                    height=0.0,
                )
            ]

        # Combine page texts
        full_text = "\n\n".join(p.text for p in pages if p.text.strip())
        page_count = len(pages)

        # Compute overall confidence
        if page_count > 0:
            overall_confidence = round(sum(p.confidence for p in pages) / float(page_count), 3)
        else:
            overall_confidence = 0.0

        # Check for unreadable pages (>30% threshold)
        unreadable_count = sum(1 for p in pages if p.confidence == 0.0 or not p.text.strip())
        if page_count > 0 and (unreadable_count / float(page_count)) > 0.3:
            warnings.append(
                f"Warning: {unreadable_count} of {page_count} pages (>30%) are unreadable or empty."
            )

        # Guess document type
        doc_type_hint = _guess_doc_type(full_text)

        return ParseResult(
            full_text=full_text,
            pages=pages,
            doc_type_hint=doc_type_hint,
            page_count=page_count,
            ocr_applied=ocr_applied,
            overall_confidence=overall_confidence,
            warnings=warnings,
            file_format=file_format,
        )

    except Exception as exc:
        logger.exception("Unexpected error while parsing document '%s': %s", file_path, exc)
        return ParseResult(
            full_text="",
            pages=[
                PageResult(
                    page_num=1,
                    text="",
                    confidence=0.0,
                    is_ocr=False,
                    width=0.0,
                    height=0.0,
                )
            ],
            doc_type_hint="unknown",
            page_count=1,
            ocr_applied=False,
            overall_confidence=0.0,
            warnings=[f"Failed to parse document: {str(exc)}"],
            file_format=Path(file_path).suffix.lstrip(".").lower() if file_path else "",
        )

    finally:
        # Clean up temporary file if created
        if temp_file_created and os.path.exists(actual_path):
            try:
                os.remove(actual_path)
            except Exception as exc:
                logger.debug("Failed to remove temporary file '%s': %s", actual_path, exc)
